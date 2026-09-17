import io
from typing import List
from jinja2 import Template
from docx import Document
from ..schemas import TranslationResponse

class ExporterService:
    def __init__(self):
        self.html_template = """
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="UTF-8">
            <style>
                body { font-family: 'Malgun Gothic', sans-serif; max-width: 800px; margin: 0 auto; padding: 20px; background-color: #f4f7f6; }
                .page-card { background: white; border-radius: 10px; padding: 20px; margin-bottom: 30px; box-shadow: 0 2px 5px rgba(0,0,0,0.1); }
                .label { font-weight: bold; color: #27ae60; font-size: 0.8em; margin-bottom: 5px; display: block; }
                .original { font-size: 1.1em; color: #2c3e50; background: #fff9c4; padding: 10px; border-radius: 5px; margin-bottom: 15px; }
                .translated { font-size: 1em; color: #34495e; padding: 10px; border-left: 4px solid #2ecc71; background: #f9f9f9; margin-bottom: 20px; }
                .vocab-table { width: 100%; border-collapse: collapse; margin-top: 10px; }
                .vocab-table th { background: #eee; padding: 8px; text-align: left; font-size: 0.9em; }
                .vocab-table td { border-bottom: 1px solid #eee; padding: 8px; }
                .word { color: #e74c3c; font-weight: bold; }
            </style>
        </head>
        <body>
            <h1>📄 학습 리포트 (ZH/EN → KO)</h1>
            {% for item in data %}
            <div class="page-card">
                <span class="label">ORIGINAL</span>
                <div class="original">{{ item.original_text }}</div>
                
                <span class="label">TRANSLATION (DeepL)</span>
                <div class="translated">{{ item.translated_text }}</div>
                
                {% if item.footer_vocab %}
                <span class="label">VOCABULARY (CHINESE ONLY)</span>
                <table class="vocab-table">
                    <thead><tr><th>단어</th><th>병음</th><th>의미</th></tr></thead>
                    <tbody>
                        {% for v in item.footer_vocab %}
                        <tr>
                            <td class="word">{{ v.word }}</td>
                            <td style="color:#3498db;">{{ v.pinyin }}</td>
                            <td>{{ v.meaning }}</td>
                        </tr>
                        {% endfor %}
                    </tbody>
                </table>
                {% endif %}
            </div>
            {% endfor %}
        </body>
        </html>
        """

    def to_html(self, data: List[TranslationResponse]) -> str:
        return Template(self.html_template).render(data=data)

    def to_word_stream(self, data: List[TranslationResponse]) -> io.BytesIO:
        doc = Document()
        for item in data:
            doc.add_paragraph(f"--- Page {item.page_number} ---").bold = True
            doc.add_paragraph(item.original_text)
            doc.add_paragraph(f"번역: {item.translated_text}")
            if item.footer_vocab:
                table = doc.add_table(rows=1, cols=3)
                table.style = 'Table Grid'
                for v in item.footer_vocab:
                    row = table.add_row().cells
                    row[0].text, row[1].text, row[2].text = v.word, v.pinyin, v.meaning
            doc.add_page_break()
        stream = io.BytesIO(); doc.save(stream); stream.seek(0)
        return stream